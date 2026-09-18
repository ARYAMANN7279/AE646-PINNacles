import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

DARK_BLUE   = RGBColor(0x1A, 0x2E, 0x4A)
MID_BLUE    = RGBColor(0x1F, 0x4E, 0x79)
ACCENT_BLUE = RGBColor(0x2E, 0x86, 0xC1)
LIGHT_BLUE  = RGBColor(0xD6, 0xE4, 0xF0)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE   = RGBColor(0xF4, 0xF6, 0xF9)
DARK_GRAY   = RGBColor(0x2C, 0x3E, 0x50)
MED_GRAY    = RGBColor(0x5D, 0x6D, 0x7E)
GREEN       = RGBColor(0x1E, 0x8B, 0x4C)
ORANGE      = RGBColor(0xCA, 0x6F, 0x1E)

W, H = Inches(13.33), Inches(7.5)
CL, CR = Inches(0.45), Inches(12.88)
CW = CR - CL
CT = Inches(1.55)
CB = Inches(7.05)
CAH = CB - CT

TOTAL = 8

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]

def slide():
    return prs.slides.add_slide(BLANK)

def box(sl, l, t, w, h, fill=None, border=None, bw=Pt(1.2)):
    s = sl.shapes.add_shape(1, l, t, w, h)
    s.line.width = bw
    if fill:  s.fill.solid(); s.fill.fore_color.rgb = fill
    else:     s.fill.background()
    if border: s.line.color.rgb = border
    else:      s.line.fill.background()
    return s

def tx(sl, text, l, t, w, h, sz=Pt(14), bold=False, color=DARK_GRAY, align=PP_ALIGN.LEFT):
    tb = sl.shapes.add_textbox(l, t, w, h)
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = "Arial"
    p.font.size = sz
    p.font.bold = bold
    p.font.color.rgb = color
    return tb

def header(sl, title, sub):
    box(sl, 0, 0, W, Inches(1.15), fill=DARK_BLUE)
    box(sl, 0, Inches(1.15), W, Inches(0.08), fill=ACCENT_BLUE)
    tx(sl, title, Inches(0.5), Inches(0.18), W-Inches(1.0), Inches(0.55),
       sz=Pt(32), bold=True, color=WHITE)
    tx(sl, sub, Inches(0.52), Inches(0.72), W-Inches(1.0), Inches(0.35),
       sz=Pt(16), color=LIGHT_BLUE)

def footer(sl, num):
    box(sl, 0, H-Inches(0.35), W, Inches(0.35), fill=DARK_BLUE)
    tx(sl, f"AE646 Stage 2 Interim Evaluation", CL, H-Inches(0.32), CW, Inches(0.3),
       sz=Pt(12), color=LIGHT_BLUE)
    tx(sl, f"{num} / {TOTAL}", W-Inches(1.5), H-Inches(0.32), Inches(1.0), Inches(0.3),
       sz=Pt(12), bold=True, color=WHITE, align=PP_ALIGN.RIGHT)

def card_body(sl, l, t, w, h):
    box(sl, l, t, w, h, fill=WHITE, border=LIGHT_BLUE, bw=Pt(1.0))

# ── S1 Title ──────────────────────────────────────────────────────────────────
sl = slide()
box(sl, 0, 0, W, H, fill=DARK_BLUE)
box(sl, 0, H-Inches(0.8), W, Inches(0.8), fill=MID_BLUE)
tx(sl, "Fourier Neural Operator for\nParametric Darcy Flow",
   Inches(1.0), Inches(2.2), W-Inches(2.0), Inches(1.5),
   sz=Pt(44), bold=True, color=WHITE)
tx(sl, "AE646 Stage 2 (Interim) Presentation",
   Inches(1.0), Inches(1.7), W-Inches(2.0), Inches(0.5),
   sz=Pt(20), color=LIGHT_BLUE)
tx(sl, "Team: PINNacles (Aryamann Srivastava, Varun Sathaye, Atishay Jain, Vedant S. Tiwari)",
   Inches(1.0), Inches(3.9), W-Inches(2.0), Inches(0.5),
   sz=Pt(16), color=LIGHT_BLUE)

# ── S2 Recap ──────────────────────────────────────────────────────────────────
sl = slide()
box(sl, 0, 0, W, H, fill=OFF_WHITE)
header(sl, "Problem Statement", "Recap of the Parametric PDE Challenge")
footer(sl, 2)
card_body(sl, CL, CT, CW, CAH)
tx(sl, "2D Darcy Flow Equation", CL+Inches(0.3), CT+Inches(0.3), CW-Inches(0.6), Inches(0.4), sz=Pt(20), bold=True, color=DARK_BLUE)
tx(sl, "-div(κ(x,y) ∇u(x,y)) = f(x,y)", CL+Inches(0.3), CT+Inches(0.8), CW-Inches(0.6), Inches(0.4), sz=Pt(18), color=DARK_GRAY)
tx(sl, "• Goal: Learn the mapping from permeability field (κ) to pressure field (u).", CL+Inches(0.3), CT+Inches(1.5), CW-Inches(0.6), Inches(0.4), sz=Pt(16))
tx(sl, "• Traditional solvers (FDM/FEM) are too slow for repeated queries.", CL+Inches(0.3), CT+Inches(2.0), CW-Inches(0.6), Inches(0.4), sz=Pt(16))
tx(sl, "• Solution: Operator Learning (surrogate modelling).", CL+Inches(0.3), CT+Inches(2.5), CW-Inches(0.6), Inches(0.4), sz=Pt(16))

# ── S3 Workflow Data ──────────────────────────────────────────────────────────
sl = slide()
box(sl, 0, 0, W, H, fill=OFF_WHITE)
header(sl, "Workflow: Data Pipeline", "Loading and Preprocessing (Completed)")
footer(sl, 3)
card_body(sl, CL, CT, CW, CAH)
tx(sl, "Dataset Generation & Loading", CL+Inches(0.3), CT+Inches(0.3), CW-Inches(0.6), Inches(0.4), sz=Pt(20), bold=True, color=DARK_BLUE)
tx(sl, "• Download: Acquired PDEBench 2D Darcy Flow (10,000 samples).", CL+Inches(0.3), CT+Inches(1.0), CW-Inches(0.6), Inches(0.4), sz=Pt(16))
tx(sl, "• Subsetting: Extracted 1,000 train/val samples and 200 test samples.", CL+Inches(0.3), CT+Inches(1.5), CW-Inches(0.6), Inches(0.4), sz=Pt(16))
tx(sl, "• Preprocessing: Downsampled train/val from 128x128 to 64x64.", CL+Inches(0.3), CT+Inches(2.0), CW-Inches(0.6), Inches(0.4), sz=Pt(16))
tx(sl, "• Spatial Features: Concatenated (x, y) coordinates to form (64, 64, 3) inputs.", CL+Inches(0.3), CT+Inches(2.5), CW-Inches(0.6), Inches(0.4), sz=Pt(16))
tx(sl, "• Normalisation: Computed statistics strictly from the training set.", CL+Inches(0.3), CT+Inches(3.0), CW-Inches(0.6), Inches(0.4), sz=Pt(16))

# ── S4 Workflow Implementation ─────────────────────────────────────────────────
sl = slide()
box(sl, 0, 0, W, H, fill=OFF_WHITE)
header(sl, "Workflow: Code & Architecture", "MLP Baseline and Initial SciML Models")
footer(sl, 4)
card_body(sl, CL, CT, CW, CAH)
tx(sl, "MLP Baseline Implementation", CL+Inches(0.3), CT+Inches(0.3), CW-Inches(0.6), Inches(0.4), sz=Pt(20), bold=True, color=DARK_BLUE)
tx(sl, "• Architecture: Flattens input to 12,288 vector -> 3x2048 hidden layers.", CL+Inches(0.3), CT+Inches(0.9), CW-Inches(0.6), Inches(0.4), sz=Pt(16))
tx(sl, "• Parameters: ~42.0 Million. No spatial inductive bias.", CL+Inches(0.3), CT+Inches(1.4), CW-Inches(0.6), Inches(0.4), sz=Pt(16))

tx(sl, "Fourier Neural Operator (SciML) Implementation", CL+Inches(0.3), CT+Inches(2.2), CW-Inches(0.6), Inches(0.4), sz=Pt(20), bold=True, color=DARK_BLUE)
tx(sl, "• Architecture: Lift to 64 channels -> 4 Spectral Convolutions (modes=12).", CL+Inches(0.3), CT+Inches(2.8), CW-Inches(0.6), Inches(0.4), sz=Pt(16))
tx(sl, "• Forward Pass: FFT -> Complex MatMul (frequency domain filtering) -> IFFT.", CL+Inches(0.3), CT+Inches(3.3), CW-Inches(0.6), Inches(0.4), sz=Pt(16))
tx(sl, "• Parameters: ~4.7 Million (Highly parameter-efficient).", CL+Inches(0.3), CT+Inches(3.8), CW-Inches(0.6), Inches(0.4), sz=Pt(16))

# ── S5 Results 1 ──────────────────────────────────────────────────────────────
sl = slide()
box(sl, 0, 0, W, H, fill=OFF_WHITE)
header(sl, "Preliminary Quantitative Results", "Relative L2 Error Comparison")
footer(sl, 5)
card_body(sl, CL, CT, CW, CAH)
tx(sl, "Model Performance (Physical Units)", CL+Inches(0.3), CT+Inches(0.3), CW-Inches(0.6), Inches(0.4), sz=Pt(20), bold=True, color=DARK_BLUE)
tx(sl, "MLP Baseline: Mean Rel L2 = 0.0820 (Params: 42.0M)", CL+Inches(0.3), CT+Inches(1.0), CW-Inches(0.6), Inches(0.4), sz=Pt(18), color=DARK_GRAY)
tx(sl, "FNO Original: Mean Rel L2 = 0.0521 (Params: 4.7M)", CL+Inches(0.3), CT+Inches(1.6), CW-Inches(0.6), Inches(0.4), sz=Pt(18), bold=True, color=GREEN)
tx(sl, "FNO Improved: Mean Rel L2 = 0.0456 (Params: 78.8M)", CL+Inches(0.3), CT+Inches(2.2), CW-Inches(0.6), Inches(0.4), sz=Pt(18), color=DARK_GRAY)

tx(sl, "Key Observations:", CL+Inches(0.3), CT+Inches(3.0), CW-Inches(0.6), Inches(0.4), sz=Pt(18), bold=True)
tx(sl, "• FNO outperforms the MLP by a significant margin using 9x fewer parameters.", CL+Inches(0.3), CT+Inches(3.6), CW-Inches(0.6), Inches(0.4), sz=Pt(16))
tx(sl, "• Error histograms show FNO errors concentrated heavily below 0.1.", CL+Inches(0.3), CT+Inches(4.1), CW-Inches(0.6), Inches(0.4), sz=Pt(16))

# ── S6 Results 2 ──────────────────────────────────────────────────────────────
sl = slide()
box(sl, 0, 0, W, H, fill=OFF_WHITE)
header(sl, "Interim Zero-Shot & Speed Benchmarks", "Super-resolution and Latency")
footer(sl, 6)
card_body(sl, CL, CT, CW, CAH)
tx(sl, "Zero-Shot Super-Resolution (128x128)", CL+Inches(0.3), CT+Inches(0.3), CW-Inches(0.6), Inches(0.4), sz=Pt(20), bold=True, color=DARK_BLUE)
tx(sl, "• Evaluated models on the native 128x128 test set (without retraining).", CL+Inches(0.3), CT+Inches(1.0), CW-Inches(0.6), Inches(0.4), sz=Pt(16))
tx(sl, "• FNO maintained high accuracy (Rel L2 = 0.0594). MLP cannot perform this task.", CL+Inches(0.3), CT+Inches(1.5), CW-Inches(0.6), Inches(0.4), sz=Pt(16))

tx(sl, "Inference Latency (The Speed Paradox)", CL+Inches(0.3), CT+Inches(2.4), CW-Inches(0.6), Inches(0.4), sz=Pt(20), bold=True, color=DARK_BLUE)
tx(sl, "• Both AI models operate in ~1 ms (1000x faster than SciPy FDM solves).", CL+Inches(0.3), CT+Inches(3.0), CW-Inches(0.6), Inches(0.4), sz=Pt(16))
tx(sl, "• Interestingly, MLP is faster per-sample (0.13ms) than FNO (0.71ms).", CL+Inches(0.3), CT+Inches(3.5), CW-Inches(0.6), Inches(0.4), sz=Pt(16))
tx(sl, "• Layer-wise profiling: spectral ops (FFT+multiply+IFFT) = 68.4% of FNO's compute;", CL+Inches(0.3), CT+Inches(4.0), CW-Inches(0.6), Inches(0.4), sz=Pt(16))
tx(sl, "  complex-multiply alone is 41.0% (the single largest component).", CL+Inches(0.3), CT+Inches(4.4), CW-Inches(0.6), Inches(0.4), sz=Pt(16))

# ── S7 Issues & Next Steps ────────────────────────────────────────────────────
sl = slide()
box(sl, 0, 0, W, H, fill=OFF_WHITE)
header(sl, "Issues Faced & Next Steps", "Moving to Stage 3")
footer(sl, 7)
card_body(sl, CL, CT, CW, CAH)
tx(sl, "Issues & Limitations Identified", CL+Inches(0.3), CT+Inches(0.3), CW-Inches(0.6), Inches(0.4), sz=Pt(20), bold=True, color=DARK_BLUE)
tx(sl, "• Issue: Near-uniform permeability samples cause model over-prediction.", CL+Inches(0.3), CT+Inches(1.0), CW-Inches(0.6), Inches(0.4), sz=Pt(16))
tx(sl, "• Cause: The dataset's Gaussian random field generation heavily favours blobby structures.", CL+Inches(0.3), CT+Inches(1.5), CW-Inches(0.6), Inches(0.4), sz=Pt(16))

tx(sl, "Next Steps for Stage 3", CL+Inches(0.3), CT+Inches(2.4), CW-Inches(0.6), Inches(0.4), sz=Pt(20), bold=True, color=DARK_BLUE)
tx(sl, "• Complete extensive hyperparameter ablations (modes, widths).", CL+Inches(0.3), CT+Inches(3.0), CW-Inches(0.6), Inches(0.4), sz=Pt(16))
tx(sl, "• Finalise quantitative plots (Efficiency Pareto Front, Error Distributions).", CL+Inches(0.3), CT+Inches(3.5), CW-Inches(0.6), Inches(0.4), sz=Pt(16))
tx(sl, "• Wrap up the Stage 3 Final Report with detailed physical interpretations.", CL+Inches(0.3), CT+Inches(4.0), CW-Inches(0.6), Inches(0.4), sz=Pt(16))

# ── S8 Summary ────────────────────────────────────────────────────────────────
sl = slide()
box(sl, 0, 0, W, H, fill=DARK_BLUE)
box(sl, 0, Inches(1.12), W, Inches(0.07), fill=ACCENT_BLUE)
tx(sl, "Summary", Inches(0.6), Inches(0.18), W-Inches(1.2), Inches(0.85), sz=Pt(36), bold=True, color=WHITE)
rows = [
    ("Status",      "Stage 2 implementation complete. Data pipeline and models validated."),
    ("Baseline",    "MLP implemented. Fast inference, but high error and no zero-shot."),
    ("SciML",       "FNO implemented. 37% lower error, fewer parameters, super-res capable."),
    ("Next Steps",  "Hyperparameter ablation and final Stage 3 report formatting.")
]
RH = Inches(1.0)
for j, (label, text) in enumerate(rows):
    RT = Inches(1.5) + j*RH
    box(sl, Inches(0.5), RT, Inches(2.0), RH-Inches(0.1), fill=ACCENT_BLUE)
    tx(sl, label, Inches(0.5), RT+Inches(0.2), Inches(2.0), RH-Inches(0.15), sz=Pt(17), bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    tx(sl, text, Inches(2.7), RT+Inches(0.25), W-Inches(3.2), RH-Inches(0.18), sz=Pt(18), color=LIGHT_BLUE)

box(sl, 0, H-Inches(0.62), W, Inches(0.62), fill=MID_BLUE)
tx(sl, "Thank you  -  Questions welcome", Inches(0.5), H-Inches(0.55), W-Inches(1.0), Inches(0.46), sz=Pt(19), color=WHITE, align=PP_ALIGN.CENTER)

import os
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "PINNacles_Stage2_Presentation.pptx")
prs.save(OUT)
print(f"Saved -> {OUT}")
