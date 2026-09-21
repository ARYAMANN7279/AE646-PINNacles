"""
Builds the Stage 3 (final) presentation: PINNacles_Stage3_FinalPresentation.pptx

Scope: optimised FNO vs tuned MLP, recipe/architecture/data/resolution studies, comparison with a
finite-volume solver, latency, physical diagnostics. Figures come from stage3/figures/
(regenerate with src/stage3_figures.py).

Run:  python3 stage3/build_stage3_presentation.py     (needs python-pptx, Pillow)
"""
import os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "figures")

DARK_BLUE   = RGBColor(0x1A, 0x2E, 0x4A)
MID_BLUE    = RGBColor(0x1F, 0x4E, 0x79)
ACCENT_BLUE = RGBColor(0x2E, 0x86, 0xC1)
LIGHT_BLUE  = RGBColor(0xD6, 0xE4, 0xF0)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE   = RGBColor(0xF4, 0xF6, 0xF9)
DARK_GRAY   = RGBColor(0x2C, 0x3E, 0x50)
GREEN       = RGBColor(0x1E, 0x8B, 0x4C)
ORANGE      = RGBColor(0xCA, 0x6F, 0x1E)

W, H = Inches(13.33), Inches(7.5)
CL, CR = Inches(0.45), Inches(12.88)
CW = CR - CL
CT = Inches(1.5)
CB = Inches(7.05)
CAH = CB - CT
TOTAL = 13

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def box(sl, l, t, w, h, fill=None, border=None, bw=Pt(1.2)):
    s = sl.shapes.add_shape(1, l, t, w, h)
    s.line.width = bw
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    else:
        s.fill.background()
    if border:
        s.line.color.rgb = border
    else:
        s.line.fill.background()
    return s


def tx(sl, text, l, t, w, h, sz=Pt(14), bold=False, color=DARK_GRAY, align=PP_ALIGN.LEFT):
    tb = sl.shapes.add_textbox(l, t, w, h)
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = "Arial"; p.font.size = sz; p.font.bold = bold; p.font.color.rgb = color
    return tb


def bullets(sl, items, l, t, w, h, sz=Pt(16), color=DARK_GRAY, gap=Pt(8)):
    tb = sl.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + item
        p.font.name = "Arial"; p.font.size = sz; p.font.color.rgb = color
        p.space_after = gap
    return tb


def pic(sl, name, l, t, maxw, maxh):
    """Insert an image scaled to fit inside (maxw x maxh), centred in that box."""
    path = os.path.join(FIG, name)
    iw, ih = Image.open(path).size
    scale = min(maxw / iw, maxh / ih)
    w, h = int(iw * scale), int(ih * scale)
    sl.shapes.add_picture(path, l + (maxw - w) // 2, t + (maxh - h) // 2, width=w, height=h)


def header(sl, title, sub):
    box(sl, 0, 0, W, H, fill=OFF_WHITE)
    box(sl, 0, 0, W, Inches(1.15), fill=DARK_BLUE)
    box(sl, 0, Inches(1.15), W, Inches(0.08), fill=ACCENT_BLUE)
    tx(sl, title, Inches(0.5), Inches(0.18), W - Inches(1.0), Inches(0.55), sz=Pt(30), bold=True, color=WHITE)
    tx(sl, sub, Inches(0.52), Inches(0.72), W - Inches(1.0), Inches(0.35), sz=Pt(16), color=LIGHT_BLUE)


def footer(sl, num):
    box(sl, 0, H - Inches(0.35), W, Inches(0.35), fill=DARK_BLUE)
    tx(sl, "AE646 Stage 3 Final Evaluation  |  Team PINNacles", CL, H - Inches(0.32), CW, Inches(0.3),
       sz=Pt(12), color=LIGHT_BLUE)
    tx(sl, f"{num} / {TOTAL}", W - Inches(1.5), H - Inches(0.32), Inches(1.0), Inches(0.3),
       sz=Pt(12), bold=True, color=WHITE, align=PP_ALIGN.RIGHT)


def card(sl, l, t, w, h, title=None, title_color=DARK_BLUE):
    box(sl, l, t, w, h, fill=WHITE, border=LIGHT_BLUE, bw=Pt(1.0))
    if title:
        tx(sl, title, l + Inches(0.25), t + Inches(0.15), w - Inches(0.5), Inches(0.4),
           sz=Pt(19), bold=True, color=title_color)



def table(sl, data, l, t, w, h, sz=Pt(14), bold_rows=(), first_left=True):
    rows, cols = len(data), len(data[0])
    tbl = sl.shapes.add_table(rows, cols, l, t, w, h).table
    for r in range(rows):
        for c in range(cols):
            cell = tbl.cell(r, c); cell.text = data[r][c]
            para = cell.text_frame.paragraphs[0]
            para.font.name = "Arial"; para.font.size = sz
            para.font.bold = (r == 0) or (r in bold_rows)
            para.alignment = PP_ALIGN.LEFT if (c == 0 and first_left) else PP_ALIGN.CENTER
            cell.fill.solid()
            cell.fill.fore_color.rgb = MID_BLUE if r == 0 else (LIGHT_BLUE if r in bold_rows else WHITE)
            para.font.color.rgb = WHITE if r == 0 else DARK_GRAY


# 1 Title
sl = slide()
box(sl, 0, 0, W, H, fill=DARK_BLUE)
box(sl, 0, H - Inches(0.8), W, Inches(0.8), fill=MID_BLUE)
tx(sl, "AE646 Stage 3 (Final) Presentation", Inches(1.0), Inches(1.5), W - Inches(2.0), Inches(0.5), sz=Pt(20), color=LIGHT_BLUE)
tx(sl, "Optimising a Fourier Neural Operator\nfor Parametric Darcy Flow", Inches(1.0), Inches(2.1), W - Inches(2.0), Inches(1.5), sz=Pt(42), bold=True, color=WHITE)
tx(sl, "Recipe, architecture, data, resolution, latency, and comparison with a finite-volume solver", Inches(1.0), Inches(3.9), W - Inches(2.0), Inches(0.5), sz=Pt(18), color=LIGHT_BLUE)
tx(sl, "Team: PINNacles (Aryamann Srivastava, Varun Sathaye, Atishay Jain, Vedant S. Tiwari)", Inches(1.0), Inches(4.7), W - Inches(2.0), Inches(0.5), sz=Pt(16), color=LIGHT_BLUE)

# 2 Recap and goals
sl = slide()
header(sl, "Recap and Stage 3 Questions", "Learning the operator kappa -> u of -div(kappa grad u) = 1, on real PDEBench data")
footer(sl, 2)
half = (CW - Inches(0.3)) / 2
card(sl, CL, CT, half, CAH, "Where Stage 2 ended", ORANGE)
bullets(sl, ["Real PDEBench pipeline: 900 / 100 / 200 split at 64x64, EDA, preprocessing and baseline ablations",
             "MLP (42 M): 0.082 mean rel. L2;  FNO (4.7 M): 0.052",
             "Speed-up claim vs an FDM solver (1030 ms)",
             ], CL + Inches(0.25), CT + Inches(0.7), half - Inches(0.5), Inches(4.5), sz=Pt(18), gap=Pt(12))
card(sl, CL + half + Inches(0.3), CT, half, CAH, "Stage 3 questions", GREEN)
bullets(sl, ["How much of the FNO error was optimisation, and how much capacity?",
             "Which architecture choices matter? How much data? Which resolution?",
             "What does inference cost, and how does the FNO compare with a proper finite-volume solver?",
             "What does the residual error mean physically?"],
        CL + half + Inches(0.55), CT + Inches(0.7), half - Inches(0.5), Inches(4.5), sz=Pt(18), gap=Pt(12))

# 3 Protocol
sl = slide()
header(sl, "Experimental Protocol", "Same split as Stage 2; every choice made on validation data only")
footer(sl, 3)
card(sl, CL, CT, CW, CAH)
bullets(sl, [
    "Data: 900 train / 100 val / 200 test real fields (seed 42); 8,800 further real samples used only for the data-scaling study",
    "Metric: per-sample relative L2 error in physical units, mean over test samples; test set evaluated once per run",
    "Reference step budget = Stage 2 (100 epochs = 5,700 steps, batch 16); seeds 0-4; sd over seeds reported",
    "Symmetry: data are D4-invariant; augmentation applied at 128x128 before sub-sampling (valid), test-time averaging over {identity, transpose}",
    "Baseline given the same effort: MLP tuned on the same footing (capacity sweep, validation selection)",
    "142 training runs on one shared GPU workstation; all configs and metrics stored per run",
], CL + Inches(0.3), CT + Inches(0.4), CW - Inches(0.6), Inches(4.6), sz=Pt(19), gap=Pt(12))

# 4 Recipe
sl = slide()
header(sl, "Finding 1: The Training Recipe Limited Stage 2", "Same models, same 5,700 steps, 3 seeds - only the recipe changes")
footer(sl, 4)
card(sl, CL, CT, Inches(7.4), CAH)
pic(sl, "fig_recipe_ablation.png", CL + Inches(0.1), CT + Inches(0.3), Inches(7.2), CAH - Inches(0.6))
card(sl, CL + Inches(7.55), CT, CW - Inches(7.55), CAH, "Result")
bullets(sl, [
    "Stage 2 recipe: MSE loss; cosine LR restarted every 200 steps",
    "Relative-L2 loss: -26%.  Single cosine: -7%.  D4 augmentation on top: -35%",
    "FNO 0.0524 -> 0.0249 (2.1x lower); MLP 0.0802 -> 0.0454",
    "Gains interact: D4 alone does not help under the old recipe",
], CL + Inches(7.75), CT + Inches(0.7), CW - Inches(7.95), Inches(4.5), sz=Pt(16), gap=Pt(10))

# 5 Architecture
sl = slide()
header(sl, "Finding 2: Architecture Choices Matter Little", "One-factor sweeps around (12 modes, width 64, 4 layers); selection on validation error")
footer(sl, 5)
card(sl, CL, CT, CW, Inches(2.6))
pic(sl, "fig_arch_sweeps.png", CL + Inches(0.1), CT + Inches(0.1), CW - Inches(0.2), Inches(2.4))
card(sl, CL, CT + Inches(2.8), CW, Inches(2.7), "What we learned", GREEN)
bullets(sl, [
    "Modes 4 to 32 (0.5 M to 34 M parameters): only 0.022-0.025 test error;  width >= 32 saturates",
    "Depth and training length help slowly: 2 -> 8 layers 0.030 -> 0.020;  5.7k -> 80k steps 0.0247 -> 0.0205",
    "Learning rate, weight decay, batch size: < 0.001 effect.  Validation picks deeper, not wider: (12, 64, 8), 9.5 M",
], CL + Inches(0.3), CT + Inches(3.4), CW - Inches(0.6), Inches(2.0), sz=Pt(16), gap=Pt(8))

# 6 Final results
sl = slide()
header(sl, "Final Models: FNO vs Tuned MLP", "60,000 steps, full recipe, 5 seeds, 200 held-out test fields")
footer(sl, 6)
table(sl, [["Model", "Params", "Mean rel. L2", "Median", "With TTA"],
           ["Stage 2 MLP", "42.0 M", "0.0820", "0.0695", "-"],
           ["Stage 2 FNO", "4.7 M", "0.0521", "0.0398", "-"],
           ["Final MLP (3x4096)", "100.7 M", "0.0415 +- 0.0009", "0.0272", "0.0386"],
           ["Final FNO (w64, L8)", "9.5 M", "0.0202 +- 0.0004", "0.0070", "0.0200"]],
      CL, CT, CW, Inches(2.6), sz=Pt(17), bold_rows=(4,))
card(sl, CL, CT + Inches(2.85), CW, Inches(2.65), "Take-aways", GREEN)
bullets(sl, [
    "FNO error 2.05x lower than an equally tuned MLP (Stage 2 gap: 1.6x) with 10.6x fewer parameters; 2.6x lower than the Stage 2 FNO",
    "Median FNO error is 0.7%: the mean is dominated by a small tail (next slides)",
    "Training the final FNO: about 32 minutes on one shared GPU",
], CL + Inches(0.3), CT + Inches(3.5), CW - Inches(0.6), Inches(2.0), sz=Pt(17), gap=Pt(8))

# 7 Data + resolution
sl = slide()
header(sl, "Finding 3: Data Scaling and Resolution", "FNO vs MLP with 100 to 9,700 samples; train-vs-evaluation resolution matrix")
footer(sl, 7)
card(sl, CL, CT, Inches(6.1), Inches(3.5))
pic(sl, "fig_data_scaling.png", CL + Inches(0.1), CT + Inches(0.1), Inches(5.9), Inches(3.3))
card(sl, CL + Inches(6.25), CT, CW - Inches(6.25), Inches(3.5))
pic(sl, "fig_resolution.png", CL + Inches(6.35), CT + Inches(0.1), CW - Inches(6.45), Inches(3.3))
card(sl, CL, CT + Inches(3.7), CW, Inches(1.8))
bullets(sl, [
    "FNO with 400 samples beats the MLP with 9,700; FNO plateaus near 0.0195 for >= 5,000 samples",
    "Zero-shot transfer 64 -> 128: 0.038 (native 128: 0.021); 64 -> 32: 0.065: useful, but 1.7-3x worse than native resolution",
], CL + Inches(0.3), CT + Inches(3.85), CW - Inches(0.6), Inches(1.6), sz=Pt(16), gap=Pt(6))

# 8 Solver
sl = slide()
header(sl, "Finding 4: Comparison with a Finite-Volume Solver", "Correction to Stage 2: the old FDM reference was inefficient - the 1443x / 7740x speed-ups are withdrawn")
footer(sl, 8)
table(sl, [["Grid", "Calibrated FV", "Error vs truth", "Direct LU", "AMG-CG"],
           ["32 x 32", "harmonic, b = 0.7", "0.094", "1.8 ms", "5.5 ms"],
           ["64 x 64", "harmonic, b = 0.8", "0.044", "7.1 ms", "11.6 ms"],
           ["128 x 128", "arithmetic, b = 1.05", "0.024", "37.4 ms", "30.7 ms"],
           ["FNO, 64 x 64 (GPU)", "9.5 M parameters", "0.020", "1.4 ms (b=1)", "0.20 ms (batched)"]],
      CL, CT, CW, Inches(2.5), sz=Pt(16), bold_rows=(4,))
card(sl, CL, CT + Inches(2.75), CW, Inches(2.75), "How to read it", ORANGE)
bullets(sl, [
    "Vectorised solver; boundary distance b calibrated on validation data (PDEBench's truth behaves as if the Dirichlet boundary were ~1 cell outside the array)",
    "At equal resolution the FNO agrees with the data 2.2x better; partly because it learns the generator's own discretisation",
    "Single CPU thread: FNO 37 ms at batch 1 - NOT faster than the 64^2 solver (7 ms); the gain is GPU batching",
], CL + Inches(0.3), CT + Inches(3.35), CW - Inches(0.6), Inches(2.1), sz=Pt(16), gap=Pt(6))

# 9 Latency
sl = slide()
header(sl, "Inference Cost: Error vs Time per Solution", "CUDA-event timing on an idle Blackwell GPU; CPU numbers from a shared host (indicative)")
footer(sl, 9)
card(sl, CL, CT, Inches(7.3), CAH)
pic(sl, "fig_pareto.png", CL + Inches(0.1), CT + Inches(0.2), Inches(7.1), CAH - Inches(0.4))
card(sl, CL + Inches(7.45), CT, CW - Inches(7.45), CAH, "Latency at 64 x 64", GREEN)
bullets(sl, [
    "FNO (9.5 M): 1.44 ms at batch 1, 0.20 ms per sample at batch 128",
    "5x (batch 1) to 35x (batched) faster than the 64^2 solver, and 26-187x faster than the 128^2 solver",
    "MLP: 0.006 ms/sample but 2x worse accuracy",
    "TF32 gives no gain; bfloat16 unsupported for complex FFT",
    "Break-even vs the 128^2 solver: ~5e4 solves",
], CL + Inches(7.65), CT + Inches(0.7), CW - Inches(7.85), Inches(4.6), sz=Pt(15), gap=Pt(8))

# 10 Physics
sl = slide()
header(sl, "Physical Interpretation", "Final FNO on the 200 test fields")
footer(sl, 10)
card(sl, CL, CT, CW, Inches(2.5))
pic(sl, "fig_diagnostics.png", CL + Inches(0.1), CT + Inches(0.1), CW - Inches(0.2), Inches(2.3))
card(sl, CL, CT + Inches(2.7), CW, Inches(2.8))
bullets(sl, [
    "Physics respected: PDE residual within 6% of the truth's own; mean pressure to 1.7%; boundary ring ratio 0.1157 vs 0.1159; 1.5e-5 of pixels negative",
    "Error is NOT concentrated at interfaces (0.0098 next to them vs 0.0075 at 8-16 cells); largest in large low-kappa regions with high pressure (r = 0.61 with peak u)",
    "Spectral: truth needs ~12 modes (1.5% truncation error), but a 4-mode FNO still reaches 2.5%: the bypass, nonlinearity and depth regenerate high frequencies",
], CL + Inches(0.3), CT + Inches(2.85), CW - Inches(0.6), Inches(2.6), sz=Pt(15), gap=Pt(6))

# 11 Samples
sl = slide()
header(sl, "Best, Median and Worst Test Sample", "kappa | truth | FNO | absolute error")
footer(sl, 11)
card(sl, CL, CT, Inches(6.6), CAH)
pic(sl, "fig_samples.png", CL + Inches(0.1), CT + Inches(0.1), Inches(6.4), CAH - Inches(0.2))
card(sl, CL + Inches(6.75), CT, CW - Inches(6.75), CAH, "Reading the figure", ORANGE)
bullets(sl, [
    "Best 0.002, median 0.007",
    "Worst 0.358: uniform kappa = 0.1; correct shape, saturated amplitude",
    "The two uniform fields have errors 0.358 and 0.085; without them the mean falls from 2.0% to 1.8%",
    "Since u ~ 1/kappa for uniform kappa, a scale-normalised formulation could remove this failure (not tested)",
], CL + Inches(6.95), CT + Inches(0.7), CW - Inches(7.15), Inches(4.6), sz=Pt(15), gap=Pt(9))

# 12 Limitations
sl = slide()
header(sl, "Limitations and Honest Notes", "What the results do and do not show")
footer(sl, 12)
card(sl, CL, CT, CW, CAH)
bullets(sl, [
    "One PDE, one dataset, binary kappa, f = 1; the plateau and scaling laws may not transfer",
    "Top architectures differ by less than seed noise; single seed per sweep point (final models: 5 seeds)",
    "Stage 2 absolute errors reflected a sub-optimal recipe (same for both models, so the comparison was fair); Stage 2 speed-up figures withdrawn",
    "Solver was calibrated to the dataset: the comparison measures agreement with PDEBench, not fidelity to the continuous PDE; single-thread, uncompiled solver",
    "Timings on a shared machine: CPU numbers are indicative; GPU timing on an idle GPU",
], CL + Inches(0.3), CT + Inches(0.4), CW - Inches(0.6), Inches(4.6), sz=Pt(19), gap=Pt(14))

# 13 Summary
sl = slide()
box(sl, 0, 0, W, H, fill=DARK_BLUE)
box(sl, 0, Inches(1.12), W, Inches(0.07), fill=ACCENT_BLUE)
tx(sl, "Summary", Inches(0.6), Inches(0.18), W - Inches(1.2), Inches(0.85), sz=Pt(36), bold=True, color=WHITE)
summary = [
    ("Recipe",   "Relative-L2 loss + single cosine + D4 augmentation halve the FNO error at equal cost (0.052 -> 0.025)"),
    ("Final",    "FNO 9.5 M: 0.0202 +- 0.0004 vs tuned MLP 100.7 M: 0.0415 +- 0.0009 (5 seeds): 2.05x, 10.6x fewer parameters"),
    ("Limits",   "Insensitive to modes/width; plateaus near 2% for >= 5,000 samples; zero-shot resolution transfer is moderate"),
    ("Solver",   "2.2x more accurate than a calibrated FV solver at 64x64 and 5-35x faster on GPU; not faster on one CPU core"),
    ("Physics",  "Respects BC, positivity, discrete PDE; residual error sits in uniform / low-contrast, high-amplitude fields"),
]
RH = Inches(0.95)
for j, (label, text) in enumerate(summary):
    RT = Inches(1.5) + j * RH
    box(sl, Inches(0.5), RT, Inches(2.0), RH - Inches(0.12), fill=ACCENT_BLUE)
    tx(sl, label, Inches(0.5), RT + Inches(0.2), Inches(2.0), RH - Inches(0.15), sz=Pt(17), bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    tx(sl, text, Inches(2.7), RT + Inches(0.12), W - Inches(3.2), RH - Inches(0.14), sz=Pt(15), color=LIGHT_BLUE)
box(sl, 0, H - Inches(0.62), W, Inches(0.62), fill=MID_BLUE)
tx(sl, "Thank you  -  Questions welcome", Inches(0.5), H - Inches(0.55), W - Inches(1.0), Inches(0.46), sz=Pt(19), color=WHITE, align=PP_ALIGN.CENTER)

OUT = os.path.join(HERE, "PINNacles_Stage3_FinalPresentation.pptx")
prs.save(OUT)
print(f"Saved -> {OUT}")
